"""
reports/excel_report.py — Writes the full fleet analysis to a formatted Excel workbook.
Requires: openpyxl
"""
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

_GB = 1024 ** 3

# Colour palette
_RED    = "FFCCCC"
_ORANGE = "FFE5CC"
_YELLOW = "FFFFCC"
_GREEN  = "CCFFCC"
_BLUE   = "CCE5FF"
_GREY   = "F2F2F2"
_WHITE  = "FFFFFF"
_DARK   = "1F3864"


def _priority_colour(priority: str) -> str:
    return {
        "REPLACE":          _RED,
        "UPGRADE_URGENT":   _ORANGE,
        "UPGRADE_PLANNED":  _YELLOW,
        "MAINTAIN":         _BLUE,
        "HEALTHY":          _GREEN,
    }.get(priority, _WHITE)


def _header(ws, row: int, cols: list[str], fill_hex: str = _DARK) -> None:
    fill = PatternFill("solid", fgColor=fill_hex)
    font = Font(bold=True, color="FFFFFF" if fill_hex == _DARK else "000000")
    for col_idx, text in enumerate(cols, 1):
        cell = ws.cell(row=row, column=col_idx, value=text)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _autofit(ws) -> None:
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)


def write_report(
    hardware_ranking: list[dict],
    cleanup_plans:    list[dict],
    sw_summary:       dict,
    od_summary:       dict,
    output_path:      str,
) -> None:
    if not HAS_OPENPYXL:
        print("openpyxl not installed — skipping Excel report. Run: pip install openpyxl")
        return

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _sheet_hardware(wb, hardware_ranking)
    _sheet_cleanup(wb, cleanup_plans)
    _sheet_software(wb, sw_summary)
    _sheet_onedrive(wb, od_summary)

    wb.save(output_path)
    print(f"Excel report saved: {output_path}")


def _sheet_hardware(wb, ranking: list[dict]) -> None:
    ws = wb.create_sheet("1. Hardware Priority")
    ws.freeze_panes = "A2"

    cols = [
        "Rank", "Hostname", "Score", "Priority", "OS", "RAM (GB)", "RAM Load %",
        "CPU Gen", "Uptime (days)", "Top Issues", "Recommended Actions",
    ]
    _header(ws, 1, cols)

    for r in ranking:
        row = [
            r["rank"],
            r["hostname"],
            r["score"],
            r["label"],
            r["os"],
            r["ram_gb"],
            r["ram_load"],
            r["cpu_gen"] if r["cpu_gen"] else "unknown",
            r["uptime_days"],
            "\n".join(r["issues"][:3]),
            "\n".join(r["actions"][:3]),
        ]
        ws.append(row)
        colour = _priority_colour(r["priority"])
        fill   = PatternFill("solid", fgColor=colour)
        for col in range(1, len(cols) + 1):
            cell = ws.cell(ws.max_row, col)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    _autofit(ws)
    ws.row_dimensions[1].height = 30
    for i in range(2, ws.max_row + 1):
        ws.row_dimensions[i].height = 45


def _sheet_cleanup(wb, plans: list[dict]) -> None:
    ws = wb.create_sheet("2. Cleanup Actions")
    ws.freeze_panes = "A2"

    cols = [
        "Hostname", "Immediate Savings", "Category", "Path", "Size", "Safe?", "Description"
    ]
    _header(ws, 1, cols)

    for plan in plans:
        hn      = plan["hostname"]
        savings = plan["immediate_savings_human"]
        for junk in plan.get("junk_cleanup", []):
            ws.append([
                hn, savings,
                junk["category"],
                junk["path"],
                junk["size_human"],
                "YES" if junk["risk"] == "none" else "REVIEW",
                junk["description"],
            ])
            row_idx = ws.max_row
            safe = junk["risk"] == "none"
            fill = PatternFill("solid", fgColor=_GREEN if safe else _YELLOW)
            for col in range(1, len(cols) + 1):
                ws.cell(row_idx, col).fill = fill

    _autofit(ws)


def _sheet_software(wb, sw_summary: dict) -> None:
    ws = wb.create_sheet("3. Software Audit")
    ws.freeze_panes = "A2"

    cols = ["Hostname", "Severity", "Software", "Issue", "Recommended Action"]
    _header(ws, 1, cols)

    sev_colour = {"high": _RED, "medium": _ORANGE, "low": _YELLOW}
    for machine in sw_summary.get("per_machine", []):
        for flag in machine.get("flags", []):
            ws.append([
                machine["hostname"],
                flag["severity"].upper(),
                flag["software"],
                flag["issue"],
                flag["action"],
            ])
            colour = sev_colour.get(flag["severity"], _WHITE)
            fill   = PatternFill("solid", fgColor=colour)
            for col in range(1, len(cols) + 1):
                ws.cell(ws.max_row, col).fill = fill

    _autofit(ws)

    # Second table: most common software
    ws2 = wb.create_sheet("3b. Common Software")
    _header(ws2, 1, ["Software Name", "Installed on N machines"])
    for item in sw_summary.get("most_common_software", []):
        ws2.append([item["name"], item["machine_count"]])
    _autofit(ws2)


def _sheet_onedrive(wb, od_summary: dict) -> None:
    ws = wb.create_sheet("4. OneDrive Migration")
    ws.freeze_panes = "A2"

    cols = [
        "Hostname", "OneDrive Ready?", "Teams Ready?",
        "Signed-in Email", "Estimated Move (GB)", "Blockers"
    ]
    _header(ws, 1, cols)

    for m in od_summary.get("per_machine", []):
        ready  = m["onedrive_ready"]
        issues = "; ".join(m.get("issues", []))
        ws.append([
            m["hostname"],
            "YES" if ready else "NO",
            "YES" if m["teams_ready"] else "NO",
            m.get("signed_in_email", ""),
            m.get("estimated_move_gb", 0),
            issues,
        ])
        fill = PatternFill("solid", fgColor=_GREEN if ready else _RED)
        for col in range(1, 3):
            ws.cell(ws.max_row, col).fill = fill

    _autofit(ws)

    # OneDrive folder structure reference
    ws3 = wb.create_sheet("4b. OneDrive Structure")
    _header(ws3, 1, ["Top-Level Folder", "Sub-folder", "Description"], fill_hex="1F3864")
    for folder, info in od_summary.get("recommended_structure", {}).items():
        ws3.append([folder, "", info.get("description", "")])
        fill = PatternFill("solid", fgColor=_BLUE)
        ws3.cell(ws3.max_row, 1).fill = fill
        ws3.cell(ws3.max_row, 1).font = Font(bold=True)
        for sub in info.get("sub", []):
            ws3.append(["", sub, ""])
    _autofit(ws3)
